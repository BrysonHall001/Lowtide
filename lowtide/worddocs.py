"""Turns tailored resume / cover letter content into Word files."""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x1B, 0x2B, 0x3C)
ACCENT = RGBColor(0x2C, 0x6A, 0x80)
MUTED = RGBColor(0x5E, 0x6E, 0x7B)
RIGHT_EDGE = Inches(7.0)


def _base():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(0)
    return d


def _run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text or "")
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", "2C6A80")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _two_sided(d, left, right, bold=False, italic=False, before=0):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_EDGE, WD_TAB_ALIGNMENT.RIGHT)
    _run(p, left, bold=bold, italic=italic, size=10.5 if bold else 10)
    if right:
        _run(p, "\t" + right, color=MUTED, size=9.5)
    return p


def resume_docx(r):
    d = _base()
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, r.get("name") or "", bold=True, size=20)
    if r.get("contact"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        _run(p, r["contact"], size=9.5, color=MUTED)
    if r.get("summary"):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _run(p, r["summary"])

    def heading(title):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        _run(p, title, bold=True, size=11, color=ACCENT)
        _rule(p)

    for sec in r.get("sections") or []:
        heading(sec.get("title") or "")
        for e in sec.get("entries") or []:
            _two_sided(d, e.get("heading") or "", e.get("dates") or "", bold=True, before=4)
            if e.get("subheading") or e.get("location"):
                _two_sided(d, e.get("subheading") or "", e.get("location") or "", italic=True)
            for b in e.get("bullets") or []:
                bp = d.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)
                _run(bp, b)

    if r.get("skills"):
        heading("Skills")
        for s in r["skills"]:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _run(p, (s.get("label") or "") + ": ", bold=True)
            _run(p, s.get("items") or "")
    return _save(d)


def cover_docx(name, text):
    d = _base()
    d.styles["Normal"].font.size = Pt(11)
    if name:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        _run(p, name, bold=True, size=16)
    for para in [x.strip() for x in (text or "").split("\n") if x.strip()]:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        _run(p, para)
    return _save(d)


def _save(d):
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf
